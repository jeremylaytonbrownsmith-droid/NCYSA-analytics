#!/usr/bin/env python3
"""
NCYSA Find My Club - GA4 data workbook builder (Phase C, Step 3).

Rebuilds the board packet .xlsx (matching NCYSA-Find-My-Club-Data.xlsx's
13-tab structure) from ga_report_data.json via openpyxl. Every number here
is the raw GA4 API export, not the GA4 UI, per the project's locked rule.

Usage: python3 generate_xlsx.py [data.json] [out.xlsx]
"""

import sys
import json
import datetime

import zipcodes
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from build_report import ZIP_LABEL_OVERRIDES

DATA_PATH = sys.argv[1] if len(sys.argv) > 1 else "ga_report_data.json"
OUT_PATH = sys.argv[2] if len(sys.argv) > 2 else "build/data_workbook.xlsx"

NAVY = "FF10045A"
GRAY = "FF666666"
WHITE = "FFFFFFFF"

TITLE_FONT = Font(name="Arial", size=16, bold=True, color=NAVY)
SUBTITLE_FONT = Font(name="Arial", size=9, color=GRAY)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color=WHITE)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(name="Arial", size=10)


MONTHS_ABBR = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun",
               "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def short_date(iso, with_year=False):
    d = datetime.date.fromisoformat(iso)
    s = f"{MONTHS_ABBR[d.month]} {d.day}"
    return f"{s}, {d.year}" if with_year else s


def zip_city(z):
    if z in ZIP_LABEL_OVERRIDES:
        return ZIP_LABEL_OVERRIDES[z]
    m = zipcodes.matching(z)
    return m[0]["city"] if m else z


def sheet_header(ws, title, subtitle, headers, col_widths, extra_subtitle=None):
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    row = 3
    if extra_subtitle:
        ws["A3"] = extra_subtitle
        ws["A3"].font = SUBTITLE_FONT
        row = 4
    header_row = row + 1
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    ws.freeze_panes = f"A{header_row + 1}"
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return header_row + 1  # first data row


def write_rows(ws, start_row, rows, percent_cols=()):
    for r_off, row in enumerate(rows):
        for c_off, val in enumerate(row, start=1):
            cell = ws.cell(row=start_row + r_off, column=c_off, value=val)
            cell.font = BODY_FONT
            if c_off in percent_cols:
                cell.number_format = "0.0%"
    return start_row + len(rows)


def sub_table(ws, row, title, headers, rows, col_offset=0):
    row += 1
    c = ws.cell(row=row, column=1 + col_offset, value=title)
    c.font = Font(name="Arial", size=10, bold=True, color=NAVY)
    row += 1
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=i + col_offset, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    row += 1
    end = write_rows_at(ws, row, rows, col_offset)
    return end


def write_rows_at(ws, start_row, rows, col_offset=0):
    for r_off, row in enumerate(rows):
        for c_off, val in enumerate(row):
            cell = ws.cell(row=start_row + r_off, column=1 + col_offset + c_off, value=val)
            cell.font = BODY_FONT
    return start_row + len(rows)


def daily_window_avg(daily, start_iso, from_date, to_date):
    start = datetime.date.fromisoformat(start_iso)
    vals = [c for i, c in enumerate(daily) if from_date <= start + datetime.timedelta(days=i) <= to_date]
    return round(sum(vals) / len(vals), 1) if vals else 0.0


def build(data, wb):
    win = data["window"]
    bc = data["baseline_check"]
    cc = data["club_click"]
    cov = data["coverage_gap"]
    fmc = data["fmc_page"]
    campaign = data["campaign"]

    zip_distinct_all = len(data["zip_counts_all"])
    zip_distinct_nc = len(data["zip_counts_nc"])
    nc_searches = sum(data["zip_counts_nc"].values())
    nc_share_pct = round(100.0 * nc_searches / sum(data["zip_counts_all"].values()))

    year = datetime.date.fromisoformat(win["start"]).year
    avg_before = daily_window_avg(data["daily"], win["start"], datetime.date(year, 3, 1), datetime.date(year, 6, 10))
    avg_after = daily_window_avg(data["daily"], win["start"], datetime.date(year, 6, 11),
                                  datetime.date.fromisoformat(win["end"]))

    period = f"{short_date(win['start'])} – {short_date(win['end'], with_year=True)}"
    click_start_label = short_date(cc["window_start"], with_year=True)

    # --- Summary ---------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "NCYSA Find My Club — GA4 Data Workbook"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Reporting period: {period} · Source: Google Analytics 4, NCYSA property"
    ws["A2"].font = SUBTITLE_FONT
    ws["A3"] = "club_click and coverage_gap tracking began Jul 8, 2026; their totals reflect ~2 weeks."
    ws["A3"].font = SUBTITLE_FONT
    row = 5
    for i, h in enumerate(["Metric", "Value", "Notes"], start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    ws.freeze_panes = "A6"
    action_total = sum(v["clicks"] for v in cc["action_detail"].values())
    website_pct = round(100.0 * cc["action_detail"].get("website", {"clicks": 0})["clicks"] / action_total) if action_total else 0
    summary_rows = [
        ("ZIP searches (zip_search)", bc["zip_search_events"],
         f"{bc['zip_search_users']:,} users, {zip_distinct_all:,} distinct ZIPs"),
        ("NC ZIP searches", nc_searches, f"{zip_distinct_nc:,} NC ZIP codes, {nc_share_pct}% of all searches"),
        ("Club clicks (club_click)", cc["total_clicks"],
         f"Since {click_start_label}; {cc['total_click_users']:,} users; {cc['distinct_clubs']:,} clubs clicked"),
        ("  → Website clicks", cc["action_detail"].get("website", {"clicks": 0})["clicks"], f"{website_pct}% of clicks"),
        ("  → Email clicks", cc["action_detail"].get("email", {"clicks": 0})["clicks"], None),
        ("  → Clicks on rank-1 club", cc["rank1_clicks"], f"{round(cc['rank1_share_pct'])}% of all clicks"),
        ("Coverage gap events", cov["total"], f"Since {click_start_label}; only {cov['nc_total']} from NC-coded ZIPs, several internal tests"),
        ("Find My Club page views", fmc["views"], f"{fmc['users']:,} users; #{fmc['rank_by_views']} page on site; {fmc['key_events']:,} key events"),
        ("worldcup2026 campaign sessions", campaign["sessions"], f"{campaign['key_events']:,} key events; traffic began {short_date(win['campaign_marker'])}"),
        ("Avg daily searches before campaign (Mar 1–Jun 10)", avg_before, None),
        (f"Avg daily searches during campaign (Jun 11–{short_date(win['end'])})", avg_after, "5x lift" if avg_before and round(avg_after / avg_before) == 5 else f"{round(avg_after / avg_before, 1)}x lift" if avg_before else None),
    ]
    write_rows(ws, row + 1, summary_rows)

    # --- Daily Searches ----------------------------------------------------
    ws = wb.create_sheet("Daily Searches")
    r0 = sheet_header(ws, "zip_search — daily event counts",
                       f"One row per day, {period}. Campaign launched mid-June.",
                       ["Date", "Searches"], [14, 12])
    start_date = win["start"]
    d0 = datetime.date.fromisoformat(start_date)
    rows = [[(d0 + datetime.timedelta(days=i)).isoformat(), c] for i, c in enumerate(data["daily"])]
    write_rows(ws, r0, rows)

    # --- Search ZIPs ---------------------------------------------------
    ws = wb.create_sheet("Search ZIPs")
    r0 = sheet_header(ws, "zip_search — by ZIP code entered",
                       f"All {zip_distinct_all:,} distinct ZIP codes searched, sorted by volume.",
                       ["ZIP", "NC ZIP?", "Searches", "Users"], [10, 10, 12, 10])
    zd = sorted(data["zip_detail"].items(), key=lambda kv: -kv[1]["searches"])
    rows = [[z, "Yes" if v["nc"] else "No", v["searches"], v["users"]] for z, v in zd]
    write_rows(ws, r0, rows)

    # --- Nearest Club (per search) ---------------------------------------
    ws = wb.create_sheet("Nearest Club (per search)")
    r0 = sheet_header(ws, "zip_search — nearest club returned",
                       "Which club appeared as the closest result, across all searches.",
                       ["Club", "Times nearest", "Users"], [40, 14, 10])
    nc = sorted(data["nearest_club"].items(), key=lambda kv: -kv[1]["events"])
    rows = [[name, v["events"], v["users"]] for name, v in nc]
    write_rows(ws, r0, rows)

    # --- Nearest Club Distance ---------------------------------------------
    ws = wb.create_sheet("Nearest Club Distance")
    r0 = sheet_header(ws, "zip_search — distance to nearest club (miles)",
                       "Each row: a distance value returned, with how many searches saw it. Sorted nearest-first.",
                       ["Miles to nearest club", "Searches", "Users"], [20, 12, 10])
    rows = [[d["miles"], d["events"], d["users"]] for d in data["nearest_club_distance"]]
    write_rows(ws, r0, rows)

    # --- Club Clicks ---------------------------------------------------
    ws = wb.create_sheet("Club Clicks")
    r0 = sheet_header(ws, "club_click — by club clicked",
                       f"Clicks to each club's website or email. Tracking began {click_start_label}.",
                       ["Club", "Clicks", "Users"], [40, 10, 10])
    cd = sorted(cc["club_detail"].items(), key=lambda kv: -kv[1]["clicks"])
    rows = [[name, v["clicks"], v["users"]] for name, v in cd]
    write_rows(ws, r0, rows)

    # --- Click Detail (two sub-tables) --------------------------------------
    ws = wb.create_sheet("Click Detail")
    ws["A1"] = "club_click — action & rank breakdowns"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Action = website vs email link. Rank = the club's position in the search results."
    ws["A2"].font = SUBTITLE_FONT
    for i, h in enumerate(["Click action", "Clicks", "Users"], start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    ad = sorted(cc["action_detail"].items(), key=lambda kv: -kv[1]["clicks"])
    row = write_rows(ws, 5, [[name, v["clicks"], v["users"]] for name, v in ad])
    row += 2
    ws.cell(row=row, column=1, value="Result rank clicked").font = Font(name="Arial", size=10, bold=True, color=NAVY)
    row += 1
    for i, h in enumerate(["Rank", "Clicks"], start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    row += 1
    write_rows(ws, row, [[int(k), v] for k, v in cc["rank_detail"].items()])
    for i, w in enumerate([16, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # --- Click ZIPs ---------------------------------------------------
    ws = wb.create_sheet("Click ZIPs")
    r0 = sheet_header(ws, "club_click — by ZIP the family searched",
                       f"Where clicking families searched from. Tracking began {click_start_label}.",
                       ["Searcher ZIP", "NC ZIP?", "Clicks", "Users"], [12, 10, 10, 10])
    czd = sorted(data["click_zips"].items(), key=lambda kv: -kv[1]["clicks"])
    rows = [[z, "Yes" if v["nc"] else "No", v["clicks"], v["users"]] for z, v in czd]
    write_rows(ws, r0, rows)

    # --- Coverage Gaps (three sub-tables) -----------------------------------
    ws = wb.create_sheet("Coverage Gaps")
    ws["A1"] = "coverage_gap — searches with no club within 40 miles"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Tracking began {click_start_label}. Most events trace to out-of-state ZIPs; NC-coded rows flagged."
    ws["A2"].font = SUBTITLE_FONT
    for i, h in enumerate(["ZIP searched", "NC ZIP?", "Events"], start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    cgz = sorted(cov["zip_detail"].items(), key=lambda kv: -kv[1]["events"])
    row = write_rows(ws, 5, [[z, "Yes" if v["nc"] else "No", v["events"]] for z, v in cgz])
    row += 2
    gap_clubs = sorted(cov["nearest_club"].items(), key=lambda kv: -kv[1]["events"])
    row = sub_table(ws, row, "Nearest club shown on gap searches", ["Club", "Events", "Users"],
                     [[name, v["events"], v["users"]] for name, v in gap_clubs])
    row += 2
    sub_table(ws, row, "Distance to nearest club on gap searches (miles)", ["Miles", "Events", "Users"],
              [[d["miles"], d["events"], d["users"]] for d in cov["nearest_club_distance"]])
    for i, w in enumerate([14, 10, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    # --- Traffic Channels ----------------------------------------------
    ws = wb.create_sheet("Traffic Channels")
    r0 = sheet_header(ws, "Sitewide sessions by channel", f"GA4 session default channel grouping, {period}.",
                       ["Channel", "Sessions", "Engaged sessions", "Engagement rate", "Key events"],
                       [18, 12, 16, 14, 12])
    ch = sorted(data["channels"].items(), key=lambda kv: -kv[1]["sessions"])
    rows = [[name, v["sessions"], v["engaged_sessions"], v["engagement_pct"] / 100, v["key_events"]] for name, v in ch]
    write_rows(ws, r0, rows, percent_cols=(4,))

    # --- First-User Channels ---------------------------------------------
    ws = wb.create_sheet("First-User Channels")
    r0 = sheet_header(ws, "New users by first-touch channel",
                       "How users first found ncsoccer.org (vs Traffic Channels tab, which counts sessions).",
                       ["Channel", "Total users", "New users", "Returning users", "Event count", "Key events"],
                       [18, 12, 12, 14, 12, 12])
    fuc = sorted(data["first_user_channels"].items(), key=lambda kv: -kv[1]["total_users"])
    rows = [[name, v["total_users"], v["new_users"], v["returning_users"], v["event_count"], v["key_events"]]
            for name, v in fuc]
    write_rows(ws, r0, rows)

    # --- Campaigns -------------------------------------------------------
    ws = wb.create_sheet("Campaigns")
    r0 = sheet_header(ws, "Sitewide sessions by campaign name",
                       "worldcup2026 is the Intrepid campaign (spans paid social + paid search).",
                       ["Campaign", "Sessions", "Engagement rate", "Key events"], [24, 12, 14, 12])
    camp = sorted(data["campaigns_full"].items(), key=lambda kv: -kv[1]["sessions"])
    rows = [[name, v["sessions"], v["engagement_pct"] / 100, v["key_events"]] for name, v in camp]
    write_rows(ws, r0, rows, percent_cols=(3,))

    # --- Top Pages (two sub-tables) --------------------------------------
    ws = wb.create_sheet("Top Pages")
    ws["A1"] = "Most viewed pages on ncsoccer.org"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Top 25 by views. Find My Club is #{fmc['rank_by_views']} by views, "
                f"#{fmc['rank_by_users']} by users, #{fmc['rank_by_key_events']} by key events.")
    ws["A2"].font = SUBTITLE_FONT
    for i, h in enumerate(["Page path", "Views", "Users", "Key events"], start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
    row = write_rows(ws, 5, [[p["path"], p["views"], p["users"], p["kev"]] for p in data["top_pages"]])
    row += 2
    sub_table(ws, row, "Find My Club page variants (all URLs)", ["Page path", "Views", "Users"],
              [[v["path"], v["views"], v["users"]] for v in data["fmc_variants"]])
    for i, w in enumerate([44, 10, 10, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    return wb


def main():
    data = json.load(open(DATA_PATH))
    wb = Workbook()
    build(data, wb)
    import os
    os.makedirs("build", exist_ok=True)
    wb.save(OUT_PATH)
    print(f"\nWrote {OUT_PATH}")


if __name__ == "__main__":
    main()
